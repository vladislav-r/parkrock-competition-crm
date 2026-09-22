import io
import uuid

from openpyxl import load_workbook
from sqlalchemy import select

from app.db import SessionLocal
from app.models import Admin, Event, RolePermission, UserRole
from app.sessions import start_session, lock_user
from app.safety_exports import create_safety_xlsx, create_safety_pdf, pages, footer_lines
from test_current_workflows import add_participant, command_headers

SETTINGS = dict(competition_name='Фестиваль «Парк Рок»', location='г. Хабаровск', dates='16–17 мая 2026 г.',
                briefing_date='16 мая 2026 г.', official_name='Иванов И.И.')


def test_safety_pages_and_xlsx():
    clubs = [dict(name='Комсомольск-на-Амуре', members=[dict(start_number=i + 1, name='Константинопольский Александр', group='Мальчики 7–9 лет') for i in range(65)]),
             dict(name='Форум Хабаровск', members=[dict(start_number=134, name='=Не формула', group='Женщины')])]
    plan = pages(clubs)
    assert len(plan) > 2
    workbook = load_workbook(io.BytesIO(create_safety_xlsx(SETTINGS, clubs)))
    assert len(workbook.worksheets) == 1
    sheet = workbook.active
    assert sheet.page_setup.orientation == 'portrait'
    assert len(sheet.row_breaks.brk) == len(plan) - 1
    borders = [0, *(b.id for b in sheet.row_breaks.brk), sheet.max_row]
    blank_rows = [row for row in sheet.iter_rows() if all(c.value is None for c in row) and all(c.border.bottom.style == 'thin' and c.border.left.style == 'thin' for c in row)]
    assert len(blank_rows) == 5 * len(clubs)
    seen = []
    for start, end, (club_name, rows, last) in zip(borders, borders[1:], plan):
        page_values = [row[0].value for row in sheet.iter_rows(min_row=start+1, max_row=end)]
        assert all((line in page_values) == last for line in footer_lines(SETTINGS))
        assert ("ИН" in [row[1].value for row in sheet.iter_rows(min_row=start+1, max_row=end)]) == (rows[0][0][0] == 1)
        assert all(sheet.row_dimensions[row[0].row].height == 12.6 for row in sheet.iter_rows(min_row=start+1, max_row=end) if isinstance(row[0].value, int))
        for row in sheet.iter_rows(min_row=start+1, max_row=end):
            if isinstance(row[0].value, int):
                assert row[3].value == club_name
                assert row[5].value == 'Б'
                assert row[6].value is None and row[7].value is None
                assert row[2].data_type == 's'
                assert row[2].font.name == 'Calibri' and row[2].font.sz == 10
                assert row[2].border.bottom.style == 'thin'
                seen.append((row[0].value, row[1].value))
    assert seen == [(i+1, i+1) for i in range(65)] + [(1, 134)]
    assert create_safety_pdf(SETTINGS, clubs).startswith(b'%PDF-')


def test_safety_permissions_and_settings(client, festival, auth_headers):
    add_participant(event_id=festival['event_id'], set_id=festival['first_set_id'], start_number=901)
    event = client.get('/api/v1/admin/exports/safety/settings', headers=auth_headers).json()
    payload = {**SETTINGS, 'expected_version': event['event_version']}
    headers = command_headers(auth_headers)
    saved = client.put('/api/v1/admin/exports/safety/settings', headers=headers, json=payload)
    assert saved.status_code == 200, saved.text
    assert client.put('/api/v1/admin/exports/safety/settings', headers=headers, json=payload).json() == saved.json()
    assert client.put('/api/v1/admin/exports/safety/settings', headers=command_headers(auth_headers), json=payload).status_code == 409
    with SessionLocal() as db:
        assert db.get(Event, festival['event_id']).export_competition_name == ''
    for role in (UserRole.reception, UserRole.secretary, UserRole.chief_judge, UserRole.route_judge):
        with SessionLocal() as db:
            user = Admin(email=f'{role.value}@safety.test', full_name='Тест', password_hash='unused', role=role, is_active=True)
            db.add(user)
            # Existing custom matrices without the new right inherit participants.manage.
            db.add_all([RolePermission(role=role, permission='exports.create', is_allowed=False),
                        RolePermission(role=role, permission='participants.manage', is_allowed=role != UserRole.route_judge)])
            db.commit()
            token = start_session(db, lock_user(db, user.id))
            db.commit()
        headers = {'Authorization': f'Bearer {token}'}
        for suffix in ('', f'?club_id={festival["club_id"]}'):
            for extension in ('pdf', 'xlsx'):
                result = client.get(f'/api/v1/admin/exports/safety.{extension}{suffix}', headers=headers)
                assert result.status_code == (403 if role == UserRole.route_judge else 200), result.text[:200]
        assert client.get('/api/v1/admin/exports/catalog', headers=headers).status_code == 403
    assert client.get(f'/api/v1/admin/exports/safety.pdf?club_id={uuid.uuid4()}', headers=auth_headers).status_code == 404
