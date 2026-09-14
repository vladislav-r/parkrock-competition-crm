import io
import uuid
from datetime import date, datetime, timezone
from fractions import Fraction

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.db import SessionLocal
from app.models import (Admin, AgeGroup, Ascent, Club, Event, EventStage, FinalCategoryResult,
    FinalRouteAttempt, Participant, QualificationCategorySnapshot, Route, Sex, UserRole)
from app.team_results import AmbiguousFsrTie, FSR_POINTS, place_points, rank_teams
from test_absolute_exports import configure
from test_stage8_judge import prepare_final, command_headers


def test_fsr_table_ties_and_exact_team_ranks():
    assert FSR_POINTS == (100,80,65,55,51,47,43,40,37,34,31,28,26,24,22,20,18,16,14,12,10,9,8,7,6,5,4,3,2,1)
    assert [place_points(i, 1) for i in range(1, 31)] == list(FSR_POINTS)
    assert place_points(31, 1) == 0
    assert place_points(2, 2) == Fraction(145, 2)
    assert place_points(29, 3) == 1
    assert place_points(30, 3) == 0
    with pytest.raises(AmbiguousFsrTie):
        place_points(30, 2)
    def member(club, place, number):
        return {"club_id": club, "place": place, "start_number": number}
    groups = [{"name": "Юноши", "rows": [member("a", 1, 1), member("b", 2, 2), member("a", 3, 3), member("a", 4, 4)]},
              {"name": "Девушки", "rows": [member("b", 1, 5), member("a", 2, 6), member("b", 3, 7)]},
              {"name": "Дети", "rows": [member(None, 1, 8), member("c", 2, 9)]}]
    rows, issues = rank_teams(groups, {"a": "Скала", "b": "Гранит", "c": "Дети"}, 2)
    assert not issues
    assert [r["points"] for r in rows] == [245, 245, 80]
    assert [r["place"] for r in rows] == [1, 1, 3]
    assert next(r for r in rows if r["club_id"] == "a")["groups"][0]["points"] == 165
    tied, _ = rank_teams([{"name": "М", "rows": [member("a", 1, 1), member("a", 2, 2), member("a", 2, 3)]}], {"a": "Скала"}, 2)
    assert tied[0]["points"] == 172.5
    assert len(tied[0]["groups"][0]["members"]) == 2
    assert rank_teams([{"name": "М", "rows": [member("a", 30, 1), member("b", 30, 2)]}], {"a": "А", "b": "Б"}, 2)[0] == []
    thirds = [{"name": str(i), "rows": [member("a", 1, 1), member(None, 1, 2), member(None, 1, 3)]} for i in range(3)]
    exact, _ = rank_teams(thirds, {"a": "А"}, 2)
    assert exact[0]["points_exact"] == "245"


def test_qualification_gate_quota_export_and_invalidation(client, festival, auth_headers):
    url = "/api/v1/public/team-results"
    assert not client.get(url).json()["available"]
    with SessionLocal() as db:
        configure(db, festival["event_id"])
        for i, routes in enumerate([festival["route_ids"], festival["route_ids"][1:], festival["route_ids"][:1]], 1):
            p = Participant(event_id=festival["event_id"], club_id=festival["club_id"], set_id=festival["first_set_id"],
                start_number=i, surname=f"Участник{i}", name="Тест", birth_date=date(2000,1,1), sex=Sex.male, club="Старое название")
            db.add(p); db.flush()
            for route in routes:
                db.add(Ascent(participant_id=p.id, route_id=route, is_completed=True))
        db.commit()
    status = client.get("/api/v1/admin/final", headers=auth_headers).json()
    confirmed = client.post("/api/v1/admin/final/qualification/confirm-all", headers=command_headers(auth_headers), json={"expected_version": status["event_version"]})
    assert confirmed.status_code == 200, confirmed.text
    result = client.get(url).json()
    assert result["available"] and result["results"][0]["points"] == 180
    assert result["results"][0]["club"] == "Тестовый клуб"
    assert not client.get(url+"?stage=final").json()["available"]
    headers = command_headers(auth_headers)
    payload = {"team_quota": 1, "expected_version": confirmed.json()["event_version"]}
    saved = client.patch("/api/v1/admin/event/team-settings", headers=headers, json=payload)
    assert saved.status_code == 200, saved.text
    assert client.patch("/api/v1/admin/event/team-settings", headers=headers, json=payload).json() == saved.json()
    assert client.patch("/api/v1/admin/event/team-settings", headers=command_headers(auth_headers), json=payload).status_code == 409
    assert client.get(url).json()["results"][0]["points"] == 100
    exported = client.get("/api/v1/admin/exports/files/teams:qualification.xlsx", headers=auth_headers)
    assert exported.status_code == 200, exported.text
    sheet = load_workbook(io.BytesIO(exported.content)).active
    assert sheet["C7"].value == 100 and sheet["D7"].value == 1
    detail = client.get("/api/v1/admin/exports/files/team-members:qualification.xlsx", headers=auth_headers)
    assert detail.status_code == 200
    assert load_workbook(io.BytesIO(detail.content)).active["G7"].value == 100
    with SessionLocal() as db:
        db.get(Route, festival["route_ids"][0]).points += 1
        db.commit()
    assert not client.get(url).json()["available"]
    assert client.get("/api/v1/admin/exports/files/teams:qualification.xlsx", headers=auth_headers).status_code == 409
    assert client.get(url+"?stage=bad").status_code == 422
    with SessionLocal() as db:
        db.get(Event, festival["event_id"]).is_public = False
        db.commit()
    assert client.get(url).status_code == 404


@pytest.mark.parametrize("quota", [0, -1, 1001, 2.5, True, "2"])
def test_quota_validation(client, festival, auth_headers, quota):
    assert client.patch("/api/v1/admin/event/team-settings", headers=command_headers(auth_headers),
        json={"team_quota": quota, "expected_version": 1}).status_code == 422


def test_permissions(client, festival, auth_headers):
    assert client.get("/api/v1/admin/team-settings").status_code == 401
    with SessionLocal() as db:
        db.get(Admin, festival["admin_id"]).role = UserRole.route_judge
        db.commit()
    assert client.get("/api/v1/admin/team-settings", headers=auth_headers).status_code == 403
    assert client.patch("/api/v1/admin/event/team-settings", headers=command_headers(auth_headers), json={"team_quota": 2, "expected_version": 1}).status_code == 403


def test_final_uses_final_places_and_qualification_snapshot(client, festival, auth_headers):
    prepare_final(client, festival, auth_headers)
    with SessionLocal() as db:
        from app.routers.admin_final import final_category_state
        event = db.get(Event, festival["event_id"])
        final = db.scalar(select(FinalCategoryResult))
        final.place = 1
        from app.models import FinalRoute
        route = db.scalar(select(FinalRoute))
        db.add(FinalRouteAttempt(event_id=event.id, final_category_result_id=final.id, final_route_id=route.id, top_attempt=1))
        db.flush()
        for group in db.scalars(select(AgeGroup)):
            category, signature = final_category_state(db, event, group.id)
            category.final_confirmed_at = datetime.now(timezone.utc)
            category.final_signature = signature
        # Live qualification changes must not affect the saved qualification places.
        for route in db.scalars(select(Route)):
            route.points = 1
        db.commit()
    qualification = client.get("/api/v1/public/team-results").json()
    final = client.get("/api/v1/public/team-results?stage=final").json()
    assert qualification["available"] and qualification["results"][0]["points"] == 100
    assert final["available"] and final["results"][0]["points"] == 100
    with SessionLocal() as db:
        attempt = db.scalar(select(FinalRouteAttempt))
        attempt.top_attempt = 2
        db.commit()
    assert not client.get("/api/v1/public/team-results?stage=final").json()["available"]
