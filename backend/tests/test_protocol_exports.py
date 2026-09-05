import io
import uuid

from openpyxl import load_workbook

from app.db import SessionLocal
from app.models import AuditLog
from sqlalchemy import select


def command_headers(auth_headers):
    return {**auth_headers, "X-Operation-Id": str(uuid.uuid4())}


def test_qualification_and_final_protocols_follow_template(client, festival, auth_headers):
    settings = client.get("/api/v1/admin/exports/settings", headers=auth_headers)
    assert settings.status_code == 200, settings.text
    saved_settings = client.put(
        "/api/v1/admin/exports/settings",
        headers=command_headers(auth_headers),
        json={
            "competition_name": "г. Хабаровск",
            "location": 'г. Хабаровск, Арена "Ерофей"',
            "dates": "16–18 октября 2026 года",
            "official_name": "Иванова А.А.",
            "official_qualification": "СС 1К",
            "expected_version": settings.json()["event_version"],
        },
    )
    assert saved_settings.status_code == 200, saved_settings.text

    participant = client.post(
        "/api/v1/admin/participants",
        headers=command_headers(auth_headers),
        json={
            "set_id": str(festival["first_set_id"]),
            "surname": "Корчак",
            "name": "Полина",
            "birth_date": "1994-05-06",
            "sex": "male",
            "sport_rank": "3",
            "club": "СШОР ЕРОФЕЙ",
            "representative": "",
            "merch_size": None,
        },
    ).json()
    participant = client.post(
        f"/api/v1/admin/participants/{participant['id']}/check-in",
        headers=command_headers(auth_headers),
        json={"expected_version": participant["version"]},
    ).json()
    result = client.put(
        f"/api/v1/admin/participants/{participant['id']}/results",
        headers=command_headers(auth_headers),
        json={
            "completed_route_ids": [str(item) for item in festival["route_ids"]],
            "expected_version": participant["version"],
        },
    )
    assert result.status_code == 200, result.text

    status = client.get("/api/v1/admin/final", headers=auth_headers).json()
    male = next(item for item in status["categories"] if item["name"] == "Мужчины")
    blocked = client.get(
        f"/api/v1/admin/exports/qualification/{male['id']}.xlsx",
        headers=auth_headers,
    )
    assert blocked.status_code == 409
    for category in status["categories"]:
        confirmed = client.post(
            f"/api/v1/admin/final/categories/{category['id']}/confirm",
            headers=command_headers(auth_headers),
            json={"expected_version": category["expected_version"]},
        )
        assert confirmed.status_code == 200, confirmed.text
        status = confirmed.json()

    qualification = client.get(
        f"/api/v1/admin/exports/qualification/{male['id']}.xlsx",
        headers=auth_headers,
    )
    assert qualification.status_code == 200, qualification.text
    assert qualification.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    workbook = load_workbook(io.BytesIO(qualification.content), data_only=True)
    sheet = workbook["Лист1"]
    assert sheet["C1"].value == "Чемпионат г. Хабаровск"
    assert sheet["C3"].value == "ИТОГОВЫЙ ПРОТОКОЛ КВАЛИФИКАЦИИ"
    assert sheet["C4"].value == "Мужчины - Боулдеринг"
    assert sheet["A5"].value == "Зам. Главного судьи по виду - Иванова А.А. (СС 1К)"
    assert [sheet[f"{column}8"].value for column in "ABCDE"] == [
        1, "Корчак Полина", "СШОР ЕРОФЕЙ", 1994, 3,
    ]
    assert sheet["F8"].value == 2
    assert sheet["H8"].value == 300
    assert all(sheet[f"{column}8"].value is None for column in "JKLMNO")

    started = client.post(
        "/api/v1/admin/final/start",
        headers=command_headers(auth_headers),
        json={"expected_version": status["event_version"]},
    )
    assert started.status_code == 200, started.text
    setup = client.get("/api/v1/admin/final/setup", headers=auth_headers).json()
    configured = client.put(
        f"/api/v1/admin/final/categories/{male['id']}/routes",
        headers=command_headers(auth_headers),
        json={
            "expected_event_version": setup["event_version"],
            "route_ids": [item["id"] for item in setup["routes"][:4]],
        },
    )
    assert configured.status_code == 200, configured.text
    final_results = client.get(
        f"/api/v1/admin/final/categories/{male['id']}/final-results",
        headers=auth_headers,
    ).json()
    final_row = final_results["results"][0]
    final_saved = client.put(
        f"/api/v1/admin/final/categories/{male['id']}/participants/{participant['id']}/final-results",
        headers=command_headers(auth_headers),
        json={
            "expected_version": final_row["version"],
            "attempts": [
                {"route_id": route["id"], "zone_attempt": 1, "top_attempt": 1}
                for route in final_results["routes"]
            ],
        },
    )
    assert final_saved.status_code == 200, final_saved.text
    status = client.get("/api/v1/admin/final", headers=auth_headers).json()
    final_confirmed = client.post(
        f"/api/v1/admin/final/categories/{male['id']}/final-confirm",
        headers=command_headers(auth_headers),
        json={"expected_version": status["event_version"]},
    )
    assert final_confirmed.status_code == 200, final_confirmed.text

    final_protocol = client.get(
        f"/api/v1/admin/exports/final/{male['id']}.xlsx",
        headers=auth_headers,
    )
    assert final_protocol.status_code == 200, final_protocol.text
    workbook = load_workbook(io.BytesIO(final_protocol.content), data_only=True)
    sheet = workbook["Лист1"]
    assert sheet["C3"].value == "ИТОГОВЫЙ ПРОТОКОЛ РЕЗУЛЬТАТОВ"
    assert [sheet[f"{column}8"].value for column in "JKLM"] == [4, 4, 4, 4]
    assert sheet["N8"].value is None
    assert sheet["O8"].value is None
    assert sheet.print_area == "'Лист1'!$A$1:$O$23"

    with SessionLocal() as db:
        actions = set(db.scalars(select(AuditLog.action)).all())
    assert {
        "export.settings.update",
        "export.qualification-protocol",
        "export.final-protocol",
    }.issubset(actions)
