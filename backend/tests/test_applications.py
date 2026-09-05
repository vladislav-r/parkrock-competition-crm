import uuid
from datetime import datetime, timezone
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from app.db import SessionLocal
from app.models import ApplicationFile, Club, Event, Participant
from app.telegram import telegram_delivery_enabled


def application_xlsx(
    surname: str = "Петров", shirt_size: str | None = "S",
    team_name: str = "Высота", representative: str = "Иванов Иван Иванович",
    sport_rank: str = "2 юношеский",
    phone: object = "79999999999", birth_year: object = 2013,
    sex: str = "М", competition_set: str = "1 (08:00–10:30)",
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Лист1"
    sheet["A4"] = "Название команды:"
    sheet["E4"] = team_name
    sheet["A5"] = "Представителем команды назначается:"
    sheet["E5"] = representative
    sheet["A6"] = "Телефон:"
    sheet["E6"] = phone
    sheet.append([])
    sheet.append([])
    sheet.append(["№ п/п", "Фамилия ", "Имя", "Отчество (можно оставить пустым)", "Год рождения", "Пол", "Разряд", "Сет", "Футболка (размер или оставить пустым)"])
    sheet.append([1, surname, "Петр", None, birth_year, sex, sport_rank, competition_set, shirt_size])
    content = BytesIO()
    workbook.save(content)
    return content.getvalue()


def operation_headers(auth_headers: dict[str, str]) -> dict[str, str]:
    return {**auth_headers, "X-Operation-Id": str(uuid.uuid4())}


def test_telegram_delivery_is_disabled_for_test_database(festival):
    assert telegram_delivery_enabled() is False


def test_public_application_can_be_downloaded_imported_and_deleted(client, festival, auth_headers):
    content = application_xlsx()
    submitted = client.post(
        "/api/v1/public/applications",
        files={"file": ("Заявка Высота.xlsm", content, "application/vnd.ms-excel.sheet.macroEnabled.12")},
    )
    assert submitted.status_code == 201, submitted.text
    application_id = submitted.json()["id"]
    assert submitted.json()["participant_count"] == 1
    assert submitted.json()["status"] == "pending"

    repeated = client.post(
        "/api/v1/public/applications",
        files={"file": ("Заявка Высота.xlsm", content, "application/vnd.ms-excel.sheet.macroEnabled.12")},
    )
    assert repeated.status_code == 201
    assert repeated.json()["id"] == application_id

    listed = client.get("/api/v1/admin/applications", headers=auth_headers)
    assert listed.status_code == 200
    assert [item["id"] for item in listed.json()] == [application_id]

    downloaded = client.get(f"/api/v1/admin/applications/{application_id}/download", headers=auth_headers)
    assert downloaded.status_code == 200
    assert downloaded.content == content

    imported = client.post(
        f"/api/v1/admin/applications/{application_id}/import",
        headers=operation_headers(auth_headers),
    )
    assert imported.status_code == 200, imported.text
    assert imported.json()["imported"] == 1

    imported_again = client.post(
        f"/api/v1/admin/applications/{application_id}/import",
        headers=operation_headers(auth_headers),
    )
    assert imported_again.status_code == 409

    with SessionLocal() as db:
        participant = db.scalar(select(Participant))
        db.delete(participant)
        db.commit()
    reimported = client.post(
        f"/api/v1/admin/applications/{application_id}/import",
        headers=operation_headers(auth_headers),
    )
    assert reimported.status_code == 200, reimported.text
    assert reimported.json()["imported"] == 1
    assert reimported.json()["import_count"] == 2
    listed_again = client.get("/api/v1/admin/applications", headers=auth_headers).json()
    assert listed_again[0]["import_count"] == 2

    deleted = client.delete(
        f"/api/v1/admin/applications/{application_id}",
        headers=operation_headers(auth_headers),
    )
    assert deleted.status_code == 200, deleted.text
    with SessionLocal() as db:
        assert db.scalar(select(func.count()).select_from(ApplicationFile)) == 0
        participant = db.scalar(select(Participant))
        assert participant is not None
        assert participant.club == "Высота"
        assert participant.merch_size is None


def test_public_application_accepts_empty_shirt_size(client, festival):
    response = client.post(
        "/api/v1/public/applications",
        files={
            "file": (
                "Заявка без футболки.xlsx",
                application_xlsx("Безфутболкин", None),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 201, response.text
    assert response.json()["participant_count"] == 1


def test_public_application_ignores_shirt_size(client, festival):
    response = client.post(
        "/api/v1/public/applications",
        files={"file": (
            "Заявка с неверным размером.xlsx",
            application_xlsx(shirt_size="4XL"),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )},
    )
    assert response.status_code == 201, response.text
    assert response.json()["participant_count"] == 1


@pytest.mark.parametrize("phone", [9999999999, 79999999999, 89999999999, "+79999999999"])
def test_public_application_accepts_supported_phone_formats(client, festival, phone):
    response = client.post(
        "/api/v1/public/applications",
        files={"file": (
            f"Заявка {phone}.xlsm",
            application_xlsx(surname=f"Телефонов{str(phone)[-2:]}", phone=phone),
            "application/vnd.ms-excel.sheet.macroEnabled.12",
        )},
    )
    assert response.status_code == 201, response.text


@pytest.mark.parametrize("phone", ["", "899 999 99 99", "8-999-999-99-99", "+89999999999", "1234567890"])
def test_public_application_rejects_invalid_phone(client, festival, phone):
    response = client.post(
        "/api/v1/public/applications",
        files={"file": (
            "Заявка с неверным телефоном.xlsm",
            application_xlsx(phone=phone),
            "application/vnd.ms-excel.sheet.macroEnabled.12",
        )},
    )
    assert response.status_code == 422, response.text
    assert "Телефон" in response.json()["detail"]


def test_public_application_rejects_invalid_representative_name(client, festival):
    response = client.post(
        "/api/v1/public/applications",
        files={"file": (
            "Заявка с неверным представителем.xlsm",
            application_xlsx(representative="Иванов И.И."),
            "application/vnd.ms-excel.sheet.macroEnabled.12",
        )},
    )
    assert response.status_code == 422, response.text
    assert "ФИО представителя" in response.json()["detail"]


@pytest.mark.parametrize(("kwargs", "field"), [
    ({"birth_year": "2013 г."}, "Год рождения"),
    ({"birth_year": datetime.now().year}, "Год рождения"),
    ({"sex": "м"}, "Пол"),
    ({"sport_rank": "Без разряда"}, "Разряд"),
    ({"competition_set": "1"}, "Сет"),
])
def test_public_application_enforces_template_values(client, festival, kwargs, field):
    response = client.post(
        "/api/v1/public/applications",
        files={"file": (
            "Заявка с неверным значением.xlsm",
            application_xlsx(**kwargs),
            "application/vnd.ms-excel.sheet.macroEnabled.12",
        )},
    )
    assert response.status_code == 422, response.text
    detail = response.json()["detail"]
    assert detail["code"] == "invalid_application"
    assert detail["rows"][0]["row_number"] == 10
    assert field in detail["rows"][0]["errors"]


def test_public_application_accepts_participant_after_row_39(client, festival):
    content = BytesIO(application_xlsx())
    workbook = load_workbook(content)
    sheet = workbook["Лист1"]
    for row_number in range(11, 41):
        sheet.cell(row_number, 2, f"Тестовый{row_number}")
        sheet.cell(row_number, 3, "Участник")
        sheet.cell(row_number, 5, 2013)
        sheet.cell(row_number, 6, "М")
        sheet.cell(row_number, 7, "б/р")
        sheet.cell(row_number, 8, "1 (08:00–10:30)")
    updated = BytesIO()
    workbook.save(updated)

    response = client.post(
        "/api/v1/public/applications",
        files={"file": (
            "Заявка на 31 участника.xlsm",
            updated.getvalue(),
            "application/vnd.ms-excel.sheet.macroEnabled.12",
        )},
    )
    assert response.status_code == 201, response.text
    assert response.json()["participant_count"] == 31


def test_public_application_finds_participant_columns_by_headers(client, festival):
    content = BytesIO(application_xlsx())
    workbook = load_workbook(content)
    sheet = workbook["Лист1"]
    sheet["B9"], sheet["C9"] = sheet["C9"].value, sheet["B9"].value
    sheet["B10"], sheet["C10"] = sheet["C10"].value, sheet["B10"].value
    updated = BytesIO()
    workbook.save(updated)

    response = client.post(
        "/api/v1/public/applications",
        files={"file": (
            "Заявка с переставленными столбцами.xlsm",
            updated.getvalue(),
            "application/vnd.ms-excel.sheet.macroEnabled.12",
        )},
    )
    assert response.status_code == 201, response.text
    assert response.json()["participant_count"] == 1


@pytest.mark.parametrize(("kwargs", "field"), [
    ({"surname": "Ф" * 101}, "Фамилия"),
    ({"team_name": "К" * 201}, "Клуб"),
    ({"representative": "П" * 201}, "Представитель"),
    ({"sport_rank": "Р" * 51}, "Разряд"),
])
def test_public_application_rejects_text_longer_than_database_limit(client, festival, kwargs, field):
    response = client.post(
        "/api/v1/public/applications",
        files={"file": (
            "Заявка с длинным полем.xlsx",
            application_xlsx(**kwargs),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )},
    )
    assert response.status_code == 422, response.text
    detail = response.json()["detail"]
    assert detail["code"] == "invalid_application"
    assert detail["rows"][0]["errors"][field].startswith("не более")
    with SessionLocal() as db:
        assert db.scalar(select(func.count()).select_from(ApplicationFile)) == 0


def test_public_application_is_duplicated_to_telegram_once(client, festival, monkeypatch):
    delivered = []

    async def capture_delivery(**application):
        delivered.append(application)

    monkeypatch.setattr("app.routers.applications.send_application_document", capture_delivery)
    content = application_xlsx("Телеграмов")

    first = client.post(
        "/api/v1/public/applications",
        files={"file": ("Заявка Телеграм.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    repeated = client.post(
        "/api/v1/public/applications",
        files={"file": ("Заявка Телеграм.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    assert first.status_code == 201, first.text
    assert repeated.status_code == 201, repeated.text
    assert delivered == [{
        "filename": "Заявка Телеграм.xlsx",
        "content": content,
        "participant_count": 1,
    }]


def test_same_club_applications_with_different_representatives_stay_separate(client, festival, auth_headers):
    application_ids = []
    for surname, representative in (("Первый", "Иванов Иван"), ("Второй", "Петров Пётр")):
        response = client.post(
            "/api/v1/public/applications",
            files={"file": (
                f"Заявка {representative}.xlsx",
                application_xlsx(surname=surname, team_name="Высота", representative=representative),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )},
        )
        assert response.status_code == 201, response.text
        application_ids.append(response.json()["id"])

    for application_id in application_ids:
        imported = client.post(
            f"/api/v1/admin/applications/{application_id}/import",
            headers=operation_headers(auth_headers),
        )
        assert imported.status_code == 200, imported.text

    clubs = client.get("/api/v1/admin/clubs", headers=auth_headers)
    assert clubs.status_code == 200, clubs.text
    height_clubs = [club for club in clubs.json() if club["name"] == "Высота"]
    assert [club["representative"] for club in height_clubs] == ["Иванов Иван", "Петров Пётр"]
    assert [club["participant_count"] for club in height_clubs] == [1, 1]
    with SessionLocal() as db:
        assert db.scalar(select(func.count()).select_from(Club).where(Club.name == "Высота")) == 2


def test_public_application_rejects_invalid_template(client, festival):
    response = client.post(
        "/api/v1/public/applications",
        files={"file": ("Заявка.xlsx", b"not an xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 422
    with SessionLocal() as db:
        assert db.scalar(select(func.count()).select_from(ApplicationFile)) == 0


def test_file_can_be_received_after_final_start_but_not_imported(client, festival, auth_headers):
    with SessionLocal() as db:
        event = db.get(Event, festival["event_id"])
        event.final_started_at = datetime.now(timezone.utc)
        db.commit()
    submitted = client.post(
        "/api/v1/public/applications",
        files={"file": ("Поздняя заявка.xlsx", application_xlsx("Поздний"), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert submitted.status_code == 201, submitted.text
    imported = client.post(
        f"/api/v1/admin/applications/{submitted.json()['id']}/import",
        headers=operation_headers(auth_headers),
    )
    assert imported.status_code == 409
    assert "запуска финала" in imported.json()["detail"]
