"""Safety register matching the approved Excel table, with explicit print pages."""
import io
import os
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.pagebreak import Break
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas

ASSETS = Path(__file__).with_name("assets")
COLUMN_WIDTHS = [3.88671875, 4.44140625, 25.5546875, 19.6640625, 15.88671875, 2.33203125, 14.109375, 14.109375]
WIDTHS = [(width * 7 + 5) * .75 for width in COLUMN_WIDTHS]
HEADERS = ["№\nп/п", "ИН", "Фамилия, имя", "Команда", "Группа", "Б", "Подпись\nинструкти-\nруемого", "Подпись\nпредставителя"]
HEADER_HEIGHTS = [48, 15, 15, 19.95, 19.95, 19.8, 43.95]
ROW_HEIGHT = 12.6
EMPTY_ROWS = 5
FOOTER_HEIGHTS = [20.4, 25.2, 15, 19.95]
PAGE_HEIGHT = 780
# Use installed Calibri on Windows; the server can supply its licensed font directory.
font_dir = Path(os.environ.get("SAFETY_FONT_DIRECTORY", "C:/Windows/Fonts"))
for name, file, fallback in [("Safety", "calibri.ttf", "DejaVuSans.ttf"), ("SafetyBold", "calibrib.ttf", "DejaVuSans-Bold.ttf")]:
    source = font_dir / file
    pdfmetrics.registerFont(TTFont(name, str(source if source.exists() else ASSETS / fallback)))


def pages(clubs):
    result = []
    for club in clubs:
        members = club["members"]
        offset = 0
        while offset < len(members):
            first = offset == 0
            available = PAGE_HEIGHT - (sum(HEADER_HEIGHTS) if first else 0)
            remaining = len(members) - offset
            last = (remaining + EMPTY_ROWS) * ROW_HEIGHT + sum(FOOTER_HEIGHTS) <= available
            capacity = remaining if last else min(int(available // ROW_HEIGHT), remaining - 1)
            chunk = []
            for index, member in enumerate(members[offset:offset + capacity], offset + 1):
                chunk.append(([index, member["start_number"], member["name"], club["name"], member["group"], "Б", "", ""], ROW_HEIGHT))
            result.append((club["name"], chunk, last))
            offset += len(chunk)
    return result


def footer_lines(settings):
    return [f'Дата проведения инструктажа: {settings.get("briefing_date") or "____________________"}',
            "Заместитель главного судьи по безопасности",
            f'____________________  ({settings.get("official_name") or "____________________________"})']


def create_safety_xlsx(settings, clubs):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Техника безопасности"
    sheet.sheet_view.showGridLines = True
    sheet.sheet_properties.pageSetUpPr.autoPageBreaks = False
    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.scale = 85
    sheet.page_margins = PageMargins(left=.25, right=.25, top=.25, bottom=.25, header=0, footer=0)
    for index, width in enumerate(COLUMN_WIDTHS, 1):
        sheet.column_dimensions[chr(64 + index)].width = width
    edge = Side(style="thin", color="000000")
    border = Border(left=edge, right=edge, top=edge, bottom=edge)
    cursor = 1

    def cell(column, value, *, bold=False, align="left", bordered=True):
        target = sheet.cell(cursor, column, value)
        if isinstance(value, str):
            target.data_type = "s"
        target.font = Font(name="Calibri", size=10, bold=bold)
        target.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
        if bordered:
            target.border = border

    def merged(value, height, *, bold=False, align="left", bordered=True, start=1, end=8):
        for col in range(start, end + 1):
            cell(col, value if col == start else "", bold=bold, align=align, bordered=bordered)
        sheet.merge_cells(start_row=cursor, start_column=start, end_row=cursor, end_column=end)
        sheet.row_dimensions[cursor].height = height

    for page_index, (club, rows, last) in enumerate(pages(clubs)):
        if page_index:
            sheet.row_breaks.append(Break(id=cursor - 1))
        first = rows[0][0][0] == 1
        if first:
            logo = Image(str(ASSETS / "climbing-federation.png"))
            logo.width, logo.height = 160, 59
            sheet.add_image(logo, f"A{cursor}")
            merged("", HEADER_HEIGHTS[0]); cursor += 1
            merged(settings["competition_name"], HEADER_HEIGHTS[1], bold=True, align="center"); cursor += 1
            merged(settings["location"], HEADER_HEIGHTS[2], start=1, end=3)
            for col in (4, 5, 6):
                cell(col, "")
            merged(settings["dates"], HEADER_HEIGHTS[2], start=7, end=8, align="right"); cursor += 1
            merged("ЖУРНАЛ", HEADER_HEIGHTS[3], bold=True, align="center"); cursor += 1
            merged("инструктажа по технике безопасности с участниками соревнований", HEADER_HEIGHTS[4], bold=True, align="center"); cursor += 1
            merged(f"Команда: {club}", HEADER_HEIGHTS[5], bold=True); cursor += 1
        table_rows = ([(HEADERS, HEADER_HEIGHTS[6])] if first else []) + rows + ([([""] * 8, ROW_HEIGHT) for _ in range(EMPTY_ROWS)] if last else [])
        for values, height in table_rows:
            for col, value in enumerate(values, 1):
                cell(col, value, bold=values is HEADERS, align="center" if col in (1, 2, 6, 7, 8) or values is HEADERS else "left")
            sheet.row_dimensions[cursor].height = height
            cursor += 1
        if last:
            merged("", FOOTER_HEIGHTS[0], bordered=False); cursor += 1
            for line, height in zip(footer_lines(settings), FOOTER_HEIGHTS[1:]):
                merged(line, height, bordered=False); cursor += 1
    sheet.print_area = f"A1:H{cursor - 1}"
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def create_safety_pdf(settings, clubs):
    output = io.BytesIO()
    canvas = Canvas(output, pagesize=(595.28, 841.89))
    canvas.setTitle("Журнал инструктажа по технике безопасности")

    def text(value, x, top, width, height, *, bold=False, align="left", bordered=False):
        canvas.saveState()
        if bordered:
            canvas.setLineWidth(.4)
            canvas.rect(x, top - height, width, height)
        clip = canvas.beginPath(); clip.rect(x + 2, top - height, width - 4, height)
        canvas.clipPath(clip, stroke=0)
        canvas.setFont("SafetyBold" if bold else "Safety", 10)
        lines = str(value).split("\n")
        baseline = top - (height - len(lines) * 12) / 2 - 9.5
        for line in lines:
            if align == "center": canvas.drawCentredString(x + width / 2, baseline, line)
            elif align == "right": canvas.drawRightString(x + width - 2, baseline, line)
            else: canvas.drawString(x + 2, baseline, line)
            baseline -= 12
        canvas.restoreState()

    for club, rows, last in pages(clubs):
        y = 823
        first = rows[0][0][0] == 1
        if first:
            text("", 20, y, sum(WIDTHS), HEADER_HEIGHTS[0], bordered=True)
            canvas.drawImage(str(ASSETS / "climbing-federation.png"), 20, y - 44, width=120, height=44, mask="auto")
            y -= HEADER_HEIGHTS[0]
            text(settings["competition_name"], 20, y, sum(WIDTHS), HEADER_HEIGHTS[1], bold=True, align="center", bordered=True); y -= HEADER_HEIGHTS[1]
            text(settings["location"], 20, y, sum(WIDTHS[:3]), HEADER_HEIGHTS[2], bordered=True)
            x = 20 + sum(WIDTHS[:3])
            for width in WIDTHS[3:6]:
                text("", x, y, width, HEADER_HEIGHTS[2], bordered=True); x += width
            text(settings["dates"], x, y, sum(WIDTHS[6:]), HEADER_HEIGHTS[2], align="right", bordered=True); y -= HEADER_HEIGHTS[2]
            for label, height, align in [("ЖУРНАЛ", HEADER_HEIGHTS[3], "center"), ("инструктажа по технике безопасности с участниками соревнований", HEADER_HEIGHTS[4], "center"), (f"Команда: {club}", HEADER_HEIGHTS[5], "left")]:
                text(label, 20, y, sum(WIDTHS), height, bold=True, align=align, bordered=True); y -= height
        table_rows = ([(HEADERS, HEADER_HEIGHTS[6])] if first else []) + rows + ([([""] * 8, ROW_HEIGHT) for _ in range(EMPTY_ROWS)] if last else [])
        for values, height in table_rows:
            x = 20
            for column, (value, width) in enumerate(zip(values, WIDTHS)):
                text(value, x, y, width, height, bold=values is HEADERS, align="center" if column in (0, 1, 5, 6, 7) or values is HEADERS else "left", bordered=True)
                x += width
            y -= height
        if last:
            y -= FOOTER_HEIGHTS[0]
            for line, height in zip(footer_lines(settings), FOOTER_HEIGHTS[1:]):
                text(line, 20, y, sum(WIDTHS), height); y -= height
        canvas.showPage()
    canvas.save()
    return output.getvalue()
