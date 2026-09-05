import csv
import io
import math
import re
import uuid
from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime

from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import func, select, text
from sqlalchemy.orm import Session

from app.clubs import get_or_create_club
from app.models import ApplicationType, CompetitionSet, Event, Participant, ParticipantSource, SetStatus, Sex
from app.services import participant_age_error


REQUIRED_COLUMNS = ("Фамилия", "Имя", "Отчество", "Дата рождения", "Пол", "Разряд", "Клуб", "Сет")
FESTIVAL_TEMPLATE_HEADERS = ("Фамилия", "Имя", "Отчество", "Год рождения", "Пол", "Разряд", "Сет")
FESTIVAL_TEMPLATE_OPTIONAL_HEADERS: tuple[str, ...] = ()
APPLICATION_RANKS = {
    "б/р", "3 юношеский", "2 юношеский", "1 юношеский",
    "3 взрослый", "2 взрослый", "1 взрослый", "КМС", "МС",
}
APPLICATION_SETS = {
    "1 (08:00–10:30)", "2 (10:45–13:15)", "3 (13:45–16:15)",
    "4 (16:30–19:00)", "5 (19:15–21:45)", "6 (08:00–10:30)",
    "7 (10:45–13:15)", "8 (13:45–16:15)", "9 (16:30–19:00)",
    "10 (19:15–21:15, 7–9 лет)",
}
PERSON_NAME_PATTERN = re.compile(r"^(?=.*[A-Za-zА-Яа-яЁё])[A-Za-zА-Яа-яЁё -]+$")
PHONE_PATTERN = re.compile(r"^(?:9\d{9}|[78]\d{10}|\+7\d{10})$")
TEXT_LENGTH_LIMITS = {
    "Фамилия": 100,
    "Имя": 100,
    "Отчество": 100,
    "Разряд": 50,
    "Клуб": 200,
    "Представитель": 200,
}


@dataclass
class ImportAnalysis:
    rows: list[dict]
    parsed_rows: list[dict]
    duplicates: int
    errors: int
    overflow: list[dict]

    def response(self) -> dict:
        valid = len(self.parsed_rows) - self.duplicates
        return {
            "total_rows": len(self.rows), "valid_rows": valid,
            "duplicate_rows": self.duplicates, "error_rows": self.errors,
            "rows": self.rows, "overflow": self.overflow,
            "can_import": self.errors == 0 and valid > 0,
        }


def read_rows(
    filename: str,
    content: bytes,
    *,
    strict_application_template: bool = False,
) -> list[dict[str, object]]:
    try:
        if filename.lower().endswith(".csv"):
            decoded = content.decode("utf-8-sig")
            lines = decoded.splitlines()
            if not lines:
                return []
            return list(csv.DictReader(io.StringIO(decoded), delimiter=";" if ";" in lines[0] else ","))
        if filename.lower().endswith((".xlsx", ".xlsm")):
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            if strict_application_template and "Лист1" not in workbook.sheetnames:
                raise HTTPException(status_code=422, detail="В шаблоне отсутствует лист «Лист1»")
            sheet = workbook["Лист1"] if strict_application_template else workbook.active
            values = list(sheet.iter_rows(values_only=True))
            if not values:
                return []
            template_rows = read_festival_template(values, strict=strict_application_template)
            if template_rows is not None:
                return template_rows
            if strict_application_template:
                raise HTTPException(status_code=422, detail="Файл не соответствует шаблону заявки")
            headers = [str(value or "").strip() for value in values[0]]
            return [dict(zip(headers, row, strict=False)) for row in values[1:] if any(value is not None for value in row)]
    except UnicodeDecodeError as exc:
        raise HTTPException(status_code=422, detail="Не удалось прочитать CSV. Сохраните файл в кодировке UTF-8.") from exc
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"Не удалось прочитать файл: {exc}") from exc
    raise HTTPException(status_code=422, detail="Поддерживаются только файлы CSV, XLSX и XLSM")


def read_festival_template(
    values: list[tuple[object, ...]],
    *,
    strict: bool = False,
) -> list[dict[str, object]] | None:
    header_row = next((index for index, row in enumerate(values) if {
        festival_template_header(value) for value in row if festival_template_header(value)
    }.issuperset(FESTIVAL_TEMPLATE_HEADERS)), None)
    if header_row is None:
        return None

    team_name = cell_text(values[3][4]) if strict and len(values) > 3 and len(values[3]) > 4 else template_field(values, "название команды")
    representative = cell_text(values[4][4]) if strict and len(values) > 4 and len(values[4]) > 4 else template_field(values, "представителем команды назначается")
    phone = cell_text(values[5][4]) if strict and len(values) > 5 and len(values[5]) > 4 else ""
    if not team_name:
        raise HTTPException(status_code=422, detail="Заполните обязательное поле «Название команды»")
    if not representative:
        raise HTTPException(status_code=422, detail="Заполните обязательное поле «Представителем команды назначается»")
    if strict and not PERSON_NAME_PATTERN.fullmatch(representative):
        raise HTTPException(
            status_code=422,
            detail="ФИО представителя может содержать только русские и латинские буквы, пробелы и дефисы",
        )
    if strict and not PHONE_PATTERN.fullmatch(phone):
        raise HTTPException(
            status_code=422,
            detail="Телефон должен иметь формат 9999999999, 79999999999, 89999999999 или +79999999999",
        )

    headers = [festival_template_header(value) for value in values[header_row]]
    rows = []
    source_rows = values[header_row + 1:]
    for excel_row_number, row in enumerate(source_rows, start=header_row + 2):
        source = dict(zip(headers, row, strict=False))
        started_columns = ("Фамилия", "Имя", "Год рождения", "Пол", "Разряд", "Сет") if strict else FESTIVAL_TEMPLATE_HEADERS
        if not any(cell_text(source.get(column)) for column in started_columns):
            continue
        rows.append({
            "Фамилия": source.get("Фамилия"),
            "Имя": source.get("Имя"),
            "Отчество": source.get("Отчество"),
            "Год рождения": source.get("Год рождения"),
            "Дата рождения": source.get("Год рождения"),
            "Пол": source.get("Пол"),
            "Разряд": source.get("Разряд"),
            "Клуб": team_name,
            "Представитель": representative,
            "Сет": source.get("Сет"),
            "__strict_application_template__": strict,
            "__row_number__": excel_row_number,
        })
    return rows


def cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value).strip()
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if math.isfinite(value) and value.is_integer():
            return format(value, ".0f")
        return str(value).strip()
    return str(value).strip()


def festival_template_header(value: object) -> str | None:
    normalized = normalize(value)
    for header in (*FESTIVAL_TEMPLATE_HEADERS, *FESTIVAL_TEMPLATE_OPTIONAL_HEADERS):
        if normalized.startswith(normalize(header)):
            return header
    return None


def template_field(values: list[tuple[object, ...]], label: str) -> str:
    normalized_label = normalize(label)
    for row in values:
        for index, value in enumerate(row):
            if normalize(value).startswith(normalized_label):
                for candidate in row[index + 1:]:
                    if str(candidate or "").strip():
                        return str(candidate).strip()
    return ""


def normalize(value: object) -> str:
    return " ".join(str(value or "").split()).casefold()


def parse_birth_date(value: object):
    if isinstance(value, datetime):
        return value.date()
    if hasattr(value, "year") and hasattr(value, "month") and hasattr(value, "day"):
        return value
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    year_text = str(value).strip()
    if isinstance(value, str):
        year_text = year_text.removesuffix("г.").removesuffix("г").strip()
    if isinstance(value, int) or (year_text.isdigit() and len(year_text) == 4):
        try:
            return date(int(year_text), 1, 1)
        except ValueError as exc:
            raise ValueError("укажите год рождения четырьмя цифрами") from exc
    for pattern in ("%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(str(value).strip(), pattern).date()
        except ValueError:
            pass
    raise ValueError("укажите дату в формате ДД.ММ.ГГГГ")


def identity(item: dict) -> tuple:
    return (
        normalize(item["surname"]), normalize(item["name"]), normalize(item["patronymic"]), item["birth_date"].year,
    )


def analyze(db: Session, event: Event, rows: list[dict[str, object]]) -> ImportAnalysis:
    if not rows:
        raise HTTPException(status_code=422, detail="Файл не содержит участников")
    if "Set" in rows[0] and "Сет" not in rows[0]:
        for row in rows:
            row["Сет"] = row.get("Set")
    missing = set(REQUIRED_COLUMNS) - set(rows[0])
    if missing:
        raise HTTPException(status_code=422, detail=f"Не хватает колонок: {', '.join(sorted(missing))}")

    competition_sets = list(db.scalars(select(CompetitionSet).where(
        CompetitionSet.event_id == event.id).order_by(CompetitionSet.name)).all())
    sets_by_name = {normalize(item.name): item for item in competition_sets}
    for item in competition_sets:
        number = "".join(char for char in item.name if char.isdigit())
        if number:
            sets_by_name[number] = item
            sets_by_name[f"сет {number}"] = item

    existing = list(db.scalars(select(Participant).where(
        Participant.event_id == event.id, Participant.archived_at.is_(None))).all())
    existing_identities = {
        (normalize(item.surname), normalize(item.name), normalize(item.patronymic), item.birth_year or item.birth_date.year)
        for item in existing
    }
    seen: set[tuple] = set()
    preview_rows: list[dict] = []
    parsed_rows: list[dict] = []
    duplicate_count = 0
    error_count = 0

    for fallback_row_number, row in enumerate(rows, start=2):
        row_number = int(row.get("__row_number__") or fallback_row_number)
        strict_application = bool(row.get("__strict_application_template__"))
        values = {
            key: "" if value is None else str(value)
            for key, value in row.items()
            if key and not key.startswith("__")
        }
        errors: dict[str, str] = {}
        birth_column = "Год рождения" if "Год рождения" in row else "Дата рождения"
        for column in REQUIRED_COLUMNS:
            source_column = birth_column if column == "Дата рождения" else column
            if column != "Отчество" and not str(row.get(column) or "").strip():
                errors[source_column] = "обязательное поле"
        for column, max_length in TEXT_LENGTH_LIMITS.items():
            if len(str(row.get(column) or "").strip()) > max_length:
                errors[column] = f"не более {max_length} символов"
        birth_date = None
        if birth_column not in errors:
            try:
                if strict_application:
                    year_text = cell_text(row.get("Дата рождения"))
                    current_year = date.today().year
                    if not re.fullmatch(r"\d{4}", year_text):
                        raise ValueError("укажите год рождения четырьмя цифрами")
                    if not current_year - 99 <= int(year_text) <= current_year - 1:
                        raise ValueError("возраст должен быть от 1 до 99 лет")
                birth_date = parse_birth_date(row.get("Дата рождения"))
                # В предпросмотре всегда показываем ровно год, даже если Excel
                # отдал ячейку как дату с нулевым временем.
                values[birth_column] = str(birth_date.year)
                if age_error := participant_age_error(birth_date, event.starts_on):
                    errors[birth_column] = age_error
            except ValueError as exc:
                errors[birth_column] = str(exc)
        sex = None
        sex_text = cell_text(row.get("Пол"))
        sex_value = normalize(sex_text)
        if strict_application and sex_text not in {"М", "Ж"}:
            errors["Пол"] = "допустимы только точные значения М или Ж"
        elif sex_value in {"м", "муж", "мужской", "male"}:
            sex = Sex.male
        elif sex_value in {"ж", "жен", "женский", "female"}:
            sex = Sex.female
        elif "Пол" not in errors:
            errors["Пол"] = "допустимо: мужской или женский"
        rank_text = cell_text(row.get("Разряд"))
        if strict_application and rank_text not in APPLICATION_RANKS and "Разряд" not in errors:
            errors["Разряд"] = "выберите разряд из списка шаблона"
        set_text = cell_text(row.get("Сет"))
        if strict_application and set_text not in APPLICATION_SETS and "Сет" not in errors:
            errors["Сет"] = "выберите сет из списка шаблона"
        set_number_match = re.match(r"^(?:сет\s*)?(\d+)\b", set_text, flags=re.IGNORECASE)
        competition_set = sets_by_name.get(normalize(set_text))
        if not competition_set and set_number_match:
            competition_set = sets_by_name.get(set_number_match.group(1))
        if not competition_set and "Сет" not in errors:
            errors["Сет"] = "сет не найден"
        elif competition_set and competition_set.status == SetStatus.confirmed:
            errors["Сет"] = "сет уже завершён"

        duplicate = False
        parsed = None
        if not errors:
            participant = {
                "surname": str(row["Фамилия"]).strip(), "name": str(row["Имя"]).strip(),
                "patronymic": str(row.get("Отчество") or "").strip(), "birth_date": birth_date,
                "birth_year": birth_date.year,
                "sex": sex, "sport_rank": str(row["Разряд"]).strip(), "club": str(row["Клуб"]).strip(),
                "representative": str(row.get("Представитель") or "").strip(),
                "merch_size": None,
            }
            item_identity = identity(participant)
            duplicate = item_identity in existing_identities or item_identity in seen
            seen.add(item_identity)
            parsed = {"row_number": row_number, "participant": participant, "competition_set": competition_set, "duplicate": duplicate}
            parsed_rows.append(parsed)
            if duplicate:
                duplicate_count += 1
        else:
            error_count += 1
        preview_rows.append({
            "row_number": row_number, "values": values, "errors": errors,
            "duplicate": duplicate, "valid": not errors and not duplicate,
        })

    current_counts = Counter(item.set_id for item in existing)
    incoming_counts = Counter(item["competition_set"].id for item in parsed_rows if not item["duplicate"])
    overflow = []
    for competition_set in competition_sets:
        incoming = incoming_counts[competition_set.id]
        projected = current_counts[competition_set.id] + incoming
        if projected > competition_set.capacity:
            overflow.append({
                "set_id": str(competition_set.id), "set_name": competition_set.name,
                "capacity": competition_set.capacity, "current": current_counts[competition_set.id],
                "incoming": incoming, "projected": projected, "overflow_by": projected - competition_set.capacity,
            })
    return ImportAnalysis(preview_rows, parsed_rows, duplicate_count, error_count, overflow)


def create_participants_from_analysis(
    db: Session,
    event: Event,
    analysis: ImportAnalysis,
    *,
    application_type: ApplicationType,
    operation_id: uuid.UUID,
) -> dict[str, int]:
    accepted = [item for item in analysis.parsed_rows if not item["duplicate"]]
    if not accepted:
        raise HTTPException(status_code=422, detail="В заявке нет корректных новых участников")
    if db.bind and db.bind.dialect.name == "postgresql":
        db.execute(text("SELECT pg_advisory_xact_lock(hashtext(:event_id))"), {"event_id": str(event.id)})
    next_number = (db.scalar(select(func.max(Participant.start_number)).where(
        Participant.event_id == event.id)) or 0) + 1
    for offset, parsed in enumerate(accepted):
        club = get_or_create_club(
            db, event.id, parsed["participant"]["club"], parsed["participant"]["representative"],
        )
        db.add(Participant(
            event_id=event.id, set_id=parsed["competition_set"].id,
            club_id=club.id, start_number=next_number + offset,
            application_type=application_type, source=ParticipantSource.import_file,
            import_operation_id=operation_id, **parsed["participant"],
        ))
    return {
        "imported": len(accepted), "skipped_duplicates": analysis.duplicates,
        "first_start_number": next_number, "last_start_number": next_number + len(accepted) - 1,
    }
